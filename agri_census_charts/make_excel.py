"""census_data.csv から、数値・出典・グラフ画像をまとめたExcelファイルを作る。

使い方: python3 make_excel.py  （先に make_charts.py でPNGを作っておく）
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).parent
FONT = "Meiryo"
HEAD_FILL = PatternFill("solid", fgColor="E8EDE3")
THIN = Side(style="thin", color="B8BEC6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
INPUT = Font(name=FONT, size=10, color="0000FF")
CALC = Font(name=FONT, size=10, color="000000")
SOURCES = [
    ("2005〜2020年", "農林水産省「農林業センサス累年統計‐農業編‐」表2-13 販売目的の作物の作物別作付（栽培）経営体数と作付（栽培）面積",
     "https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&lid=000001407063"),
    ("2025年", "農林水産省「2025年農林業センサス（確定値）第2巻 総括編」表Ⅱ-13 水稲（食用）／表Ⅱ-16 大豆",
     "https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&lid=000001487916"),
]
NOTES = [
    "全国・農業経営体ベース、販売目的で作付けした面積と経営体数。",
    "センサスの集計単位は「農家戸数」ではなく「農業経営体」。",
    "平均作付面積（ha／経営体）＝作付面積÷作付経営体数（シート内の数式で計算）。",
    "作物統計（毎年の作付面積統計）とは対象範囲が異なるため、数値は一致しない。",
    "青字＝出典からの転記値、黒字＝数式による計算値。",
]


def load():
    with open(HERE / "census_data.csv", encoding="utf-8") as f:
        rows = list(csv.DictReader(line for line in f if not line.startswith("#")))
    by_crop = {}
    for r in rows:
        by_crop.setdefault(r["crop"], []).append(r)
    return by_crop


def data_sheet(ws, crop, rows):
    ws["A1"] = f"{crop}：作付面積と1経営体当たり平均作付面積（全国）"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    heads = ["年", "作付面積（ha）", "作付面積（千ha）", "作付経営体数（経営体）", "平均作付面積（ha／経営体）", "出典"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.fill, cell.border = HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(rows, 4):
        vals = [
            (f"{r['year']}年", INPUT, "@"),
            (int(r["area_ha"]), INPUT, "#,##0"),
            (f"=B{i}/1000", CALC, "#,##0.0"),
            (int(r["entities"]), INPUT, "#,##0"),
            (f"=IFERROR(B{i}/D{i},\"\")", CALC, "0.00"),
            (r["source"], CALC, "@"),
        ]
        for c, (v, font, fmt) in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.number_format, cell.border = font, fmt, BOX
            if c == 1:
                cell.alignment = Alignment(horizontal="center")
    last = 3 + len(rows)
    ws.cell(row=last + 2, column=1, value="※ 青字＝出典からの転記値、黒字＝数式による計算値。出典の詳細は「出典・注記」シート。").font = Font(name=FONT, size=9, color="5B6470")
    for col, w in zip("ABCDEF", (10, 16, 16, 22, 26, 44)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "A4"
    img = Image(HERE / f"chart_{'suito' if crop == '水稲' else 'daizu'}.png")
    img.width, img.height = 900, 540
    ws.add_image(img, f"A{last + 4}")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for crop, rows in load().items():
        data_sheet(wb.create_sheet(crop), crop, rows)
    ws = wb.create_sheet("出典・注記")
    ws["A1"] = "出典"
    ws["A1"].font = Font(name=FONT, size=12, bold=True)
    for i, (yr, name, url) in enumerate(SOURCES, 2):
        ws.cell(row=i, column=1, value=yr).font = CALC
        ws.cell(row=i, column=2, value=name).font = CALC
        link = ws.cell(row=i, column=3, value=url)
        link.hyperlink, link.font = url, Font(name=FONT, size=10, color="0563C1", underline="single")
    n = 3 + len(SOURCES)
    ws.cell(row=n, column=1, value="注記").font = Font(name=FONT, size=12, bold=True)
    for j, note in enumerate(NOTES, n + 1):
        ws.cell(row=j, column=1, value=f"・{note}").font = CALC
    for col, w in zip("ABC", (14, 90, 60)):
        ws.column_dimensions[col].width = w
    out = HERE / "農林業センサス_水稲大豆_作付面積.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
