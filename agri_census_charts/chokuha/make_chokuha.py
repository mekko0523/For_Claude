"""水稲の直播栽培に関するグラフ（PNG4枚）とExcelを作る。

使い方: python3 make_chokuha.py
"""
import csv
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).parent
KANDEN, TANSUI = "#B5763A", "#4A90C2"
ISHOKU, CHOKU = "#8A9199", "#2B4C7E"
UP, DOWN = "#5E9E4E", "#B8BEC6"
INK, MUTED, GRID = "#1F2328", "#5B6470", "#E3E6EA"
SRC_MAFF = "資料：農林水産省調べ（都道府県への任意の聞き取り調査。未回答の県があり、統計としての精度は確保されていない）"


def set_font():
    for name in ("IPAGothic", "IPAexGothic", "Noto Sans CJK JP"):
        if any(f.name == name for f in font_manager.fontManager.ttflist):
            plt.rcParams["font.family"] = name
            return
    sys.exit("日本語フォントが見つかりません")


def read(name):
    with open(HERE / name, encoding="utf-8") as f:
        return list(csv.DictReader(line for line in f if not line.startswith("#")))


def style(ax, title, subtitle=None):
    ax.set_title(title, loc="left", fontsize=14, color=INK, pad=34 if subtitle else 14)
    if subtitle:
        ax.text(0, 1.02, subtitle, transform=ax.transAxes, fontsize=9.5, color=MUTED, va="bottom")
    ax.tick_params(colors=MUTED, length=0)
    for s in ax.spines.values():
        s.set_visible(False)


def footer(fig, *lines):
    fig.text(0.01, 0.01, "\n".join(lines), fontsize=8, color=MUTED, va="bottom")


def chart_trend(rows, out):
    years = [int(r["year"]) for r in rows]
    kan = [float(r["kanden_ha"]) / 1000 for r in rows]
    tan = [float(r["tansui_ha"]) / 1000 for r in rows]
    tot = [float(r["total_ha"]) / 1000 for r in rows]
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=200)
    ax.bar(years, tan, width=0.72, color=TANSUI, label="湛水直播", zorder=2)
    ax.bar(years, kan, width=0.72, bottom=tan, color=KANDEN, label="乾田直播", zorder=2)
    for y, t in zip(years, tot):
        if y in (1995, 2005, 2015, 2024):
            ax.annotate(f"{t:,.1f}", (y, t), xytext=(0, 4), textcoords="offset points",
                        ha="center", fontsize=10, color=INK, fontweight="bold")
    ax.axvspan(2018.5, 2024.5, color="#F2F4F6", zorder=0)
    ax.text(2021.5, max(tot) * 1.13, "公表値", ha="center", fontsize=8.5, color=MUTED)
    ax.text(2006.5, max(tot) * 1.13, "農水省資料のグラフからの読取値", ha="center", fontsize=8.5, color=MUTED)
    ax.set_xticks([y for y in years if y % 5 == 0 or y == 2024], [f"{y}年" for y in years if y % 5 == 0 or y == 2024])
    ax.set_ylabel("直播栽培面積（千ha）", color=INK)
    ax.set_ylim(0, max(tot) * 1.2)
    ax.set_xlim(1994.3, 2024.7)
    ax.grid(axis="y", color=GRID, zorder=0)
    style(ax, "水稲の直播栽培面積の推移（全国）",
          f"1995年の{tot[0]:.1f}千haから2024年は{tot[-1]:.1f}千haへ約{tot[-1] / tot[0]:.1f}倍。近年は乾田直播が伸び、2023年に湛水直播を逆転")
    ax.legend(loc="upper left", frameon=False, fontsize=9.5, labelcolor=INK)
    footer(fig, SRC_MAFF, "2018年以前は農林水産省「乾田直播・湛水直播の現状について」（令和7年12月）のグラフから換算（公表値との差±25ha以内、10ha単位に丸め）")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    fig.savefig(out, facecolor="white")
    plt.close(fig)


def chart_index(rows, out):
    rows = [r for r in rows if r["kind"] == "公表値"]
    years = [int(r["year"]) for r in rows]
    ishoku = [float(r["suito_ha"]) - float(r["total_ha"]) for r in rows]
    choku = [float(r["total_ha"]) for r in rows]
    kan = [float(r["kanden_ha"]) for r in rows]
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=200)
    series = [("移植栽培", ishoku, ISHOKU), ("直播栽培（合計）", choku, CHOKU), ("うち乾田直播", kan, KANDEN)]
    for name, vals, color in series:
        idx = [v / vals[0] * 100 for v in vals]
        ax.plot(years, idx, color=color, linewidth=2.4, marker="o", markersize=7,
                markeredgecolor="white", markeredgewidth=1.5, zorder=3)
        ax.annotate(f"{name}  {idx[-1]:.0f}", (years[-1], idx[-1]), xytext=(10, 0), textcoords="offset points",
                    va="center", fontsize=10.5, color=INK, fontweight="bold")
    ax.axhline(100, color="#B8BEC6", linewidth=1, zorder=1)
    ax.set_xticks(years, [f"{y}年" for y in years])
    ax.set_xlim(years[0] - 0.3, years[-1] + 1.9)
    ax.set_ylabel("指数（2019年＝100）", color=INK)
    ax.grid(axis="y", color=GRID, zorder=0)
    share = [c / float(r["suito_ha"]) * 100 for c, r in zip(choku, rows)]
    style(ax, "移植栽培と直播栽培の面積の推移（全国、2019年＝100）",
          f"移植は{ishoku[0] / 1000:,.0f}千ha→{ishoku[-1] / 1000:,.0f}千haに減少、直播は{choku[0] / 1000:.1f}千ha→{choku[-1] / 1000:.1f}千haに増加"
          f"（直播の普及率 {share[0]:.1f}%→{share[-1]:.1f}%）")
    footer(fig, SRC_MAFF + "「都道府県別、乾田湛水直播別水稲栽培面積」各年産",
           "移植栽培面積＝水稲作付面積－直播栽培面積（水稲作付面積は飼料用米・WCS用稲等を含む）")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    fig.savefig(out, facecolor="white")
    plt.close(fig)


def chart_region(rows, out):
    nat = 41656.7 / 1514380 * 100
    data = [(r["region"], float(r["chokuha_ha"]) / float(r["suito_ha"]) * 100, float(r["chokuha_ha"])) for r in rows]
    data = sorted(data + [("全国", nat, 41656.7)], key=lambda d: d[1])
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=200)
    y = range(len(data))
    ax.barh(y, [d[1] for d in data], height=0.6, zorder=2,
            color=[CHOKU if d[0] == "全国" else TANSUI for d in data])
    for i, (name, pct, ha) in enumerate(data):
        ax.text(pct + 0.12, i, f"{pct:.1f}%（{ha / 1000:,.1f}千ha）", va="center", fontsize=10, color=INK,
                fontweight="bold" if name == "全国" else "normal")
    ax.set_yticks(list(y), [d[0] for d in data], color=INK, fontsize=10.5)
    ax.get_yticklabels()[[d[0] for d in data].index("全国")].set_fontweight("bold")
    ax.set_xlim(0, max(d[1] for d in data) * 1.35)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    ax.set_xlabel("直播栽培の普及率（直播栽培面積÷水稲作付面積）", color=INK)
    ax.grid(axis="x", color=GRID, zorder=0)
    style(ax, "地域別の水稲直播栽培の普及率（2024年産）",
          "東海・北海道・北陸で全国平均を大きく上回る。面積では東北が最大")
    footer(fig, SRC_MAFF + "「都道府県別、乾田湛水直播別水稲栽培面積（令和6年産）」", "沖縄（直播0ha）は地域の棒から省略（全国には含む）")
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(out, facecolor="white")
    plt.close(fig)


def chart_size(rows, out):
    labels = [r["size"] for r in rows]
    chg = [(int(r["y2025"]) / int(r["y2020"]) - 1) * 100 for r in rows]
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=200)
    x = range(len(rows))
    ax.bar(x, chg, width=0.6, color=[UP if c > 0 else DOWN for c in chg], zorder=2)
    for i, (c, r) in enumerate(zip(chg, rows)):
        ax.annotate(f"{c:+.1f}%", (i, c), xytext=(0, 6 if c > 0 else -6), textcoords="offset points",
                    ha="center", va="bottom" if c > 0 else "top", fontsize=11, color=INK, fontweight="bold")
        ax.text(i, -36, f"{int(r['y2020']):,}\n→{int(r['y2025']):,}", ha="center", va="top", fontsize=8.5, color=MUTED)
    ax.axhline(0, color="#8A9199", linewidth=1)
    ax.set_xticks(list(x), labels, color=INK, fontsize=10.5)
    ax.set_ylim(-46, 24)
    ax.set_yticks(range(-30, 21, 10))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:+.0f}%" if v else "0%"))
    ax.set_ylabel("経営体数の増減率（2020年→2025年）", color=INK)
    ax.grid(axis="y", color=GRID, zorder=0)
    tot20 = sum(int(r["y2020"]) for r in rows)
    tot25 = sum(int(r["y2025"]) for r in rows)
    style(ax, "水稲の作付面積規模別経営体数の増減（全国、2020年→2025年）",
          f"経営体全体は{(tot25 / tot20 - 1) * 100:+.1f}%と大きく減る一方、10ha以上の層だけが増加し、担い手への集約が進む")
    footer(fig, "資料：農林水産省「2025年農林業センサス（確定値）」第2巻 総括編 表Ⅱ-14（販売目的の水稲（食用）作付面積規模別経営体数）",
           "棒の下の数字は経営体数（2020年→2025年）")
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(out, facecolor="white")
    plt.close(fig)


# ---- Excel ----
FONT = "Meiryo"
HEAD_FILL = PatternFill("solid", fgColor="E8EDE3")
THIN = Side(style="thin", color="B8BEC6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")


def table(ws, title, note, heads, body, widths, image):
    """body: 行ごとに [(値, 青字か, 表示形式), ...]。数式は文字列 '=...' で渡す。"""
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, size=9, color="5B6470")
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.border = Font(name=FONT, size=10, bold=True), HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 34
    for i, row in enumerate(body, 5):
        for c, (v, is_input, fmt) in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font, cell.number_format, cell.border = (BLUE if is_input else BLACK), fmt, BOX
    last = 4 + len(body)
    ws.cell(row=last + 1, column=1, value="※ 青字＝出典からの転記値、黒字＝数式による計算値。出典の詳細は「出典・注記」シート。").font = Font(name=FONT, size=9, color="5B6470")
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    img = Image(HERE / image)
    img.width, img.height = 900, 504
    ws.add_image(img, f"A{last + 3}")
    return last


def build_excel(area, region, size):
    wb = Workbook()
    wb.remove(wb.active)

    body = []
    for i, r in enumerate(area, 5):
        body.append([(int(r["year"]), True, "0"), (r["label"], True, "@"),
                     (float(r["kanden_ha"]), True, "#,##0"), (float(r["tansui_ha"]), True, "#,##0"),
                     (float(r["total_ha"]), True, "#,##0"), (f"=IFERROR(C{i}/E{i},\"\")", False, "0.0%"),
                     (r["kind"], True, "@")])
    table(wb.create_sheet("直播面積推移"), "水稲の直播栽培面積の推移（全国、1995〜2024年）",
          "2018年以前は農水省資料のグラフからの読取値（公表値との差±25ha以内）、2019年以降は公表値",
          ["年", "和暦", "乾田直播（ha）", "湛水直播（ha）", "直播合計（ha）", "直播に占める乾田直播の割合", "区分"],
          body, (8, 11, 15, 15, 15, 16, 10), "chart_chokuha_trend.png")

    pub = [r for r in area if r["kind"] == "公表値"]
    body = []
    for i, r in enumerate(pub, 5):
        body.append([(int(r["year"]), True, "0"), (float(r["suito_ha"]), True, "#,##0"),
                     (float(r["total_ha"]), True, "#,##0.0"), (float(r["kanden_ha"]), True, "#,##0.0"),
                     (f"=B{i}-C{i}", False, "#,##0"), (f"=C{i}/B{i}", False, "0.00%"),
                     (f"=E{i}/$E$5*100", False, "0.0"), (f"=C{i}/$C$5*100", False, "0.0"), (f"=D{i}/$D$5*100", False, "0.0")])
    table(wb.create_sheet("移植と直播"), "移植栽培と直播栽培の面積の推移（全国、2019〜2024年）",
          "移植栽培面積＝水稲作付面積－直播栽培面積。水稲作付面積は飼料用米・WCS用稲等を含む（農水省調査の値）",
          ["年", "水稲作付面積（ha）", "直播栽培面積（ha）", "うち乾田直播（ha）", "移植栽培面積（ha）",
           "直播の普及率", "移植 指数（2019=100）", "直播 指数（2019=100）", "乾田直播 指数（2019=100）"],
          body, (8, 16, 16, 16, 16, 12, 14, 14, 16), "chart_ishoku_vs_chokuha.png")

    body = []
    for i, r in enumerate(region, 5):
        body.append([(r["region"], True, "@"), (float(r["suito_ha"]), True, "#,##0"),
                     (float(r["chokuha_ha"]), True, "#,##0.0"), (f"=C{i}/B{i}", False, "0.0%")])
    n = 5 + len(region)
    body.append([("全国（沖縄含む）", True, "@"), (1514380, True, "#,##0"), (41656.7, True, "#,##0.0"), (f"=C{n}/B{n}", False, "0.0%")])
    table(wb.create_sheet("地域別普及率"), "地域別の水稲直播栽培の普及率（2024年産）",
          "普及率＝直播栽培面積÷水稲作付面積（水稲作付面積は飼料用米・WCS用稲等を含む）",
          ["地域", "水稲作付面積（ha）", "直播栽培面積（ha）", "普及率"], body, (18, 18, 18, 10), "chart_chokuha_region.png")

    body = []
    for i, r in enumerate(size, 5):
        body.append([(r["size"], True, "@"), (int(r["y2020"]), True, "#,##0"), (int(r["y2025"]), True, "#,##0"),
                     (f"=C{i}-B{i}", False, "+#,##0;-#,##0;0"), (f"=C{i}/B{i}-1", False, "+0.0%;-0.0%;0.0%")])
    n = 5 + len(size)
    body.append([("合計", False, "@"), (f"=SUM(B5:B{n - 1})", False, "#,##0"), (f"=SUM(C5:C{n - 1})", False, "#,##0"),
                 (f"=C{n}-B{n}", False, "+#,##0;-#,##0;0"), (f"=C{n}/B{n}-1", False, "+0.0%;-0.0%;0.0%")])
    table(wb.create_sheet("規模別経営体数"), "販売目的の水稲（食用）作付面積規模別経営体数（全国）",
          "原表の階層（0.1ha未満〜15ha以上の12区分）を6区分に合算",
          ["作付面積規模", "2020年（経営体）", "2025年（経営体）", "増減（経営体）", "増減率"],
          body, (16, 16, 16, 14, 10), "chart_suito_size.png")

    ws = wb.create_sheet("出典・注記")
    lines = [
        ("出典", None),
        ("農林水産省「都道府県別、乾田湛水直播別水稲栽培面積」令和元年産〜令和6年産（令和8年6月データ訂正）",
         "https://www.maff.go.jp/j/syouan/keikaku/soukatu/chokuha.html"),
        ("農林水産省「乾田直播・湛水直播の現状について」（令和7年12月）p.3「水稲直播栽培面積の推移」",
         "https://www.maff.go.jp/j/nousin/attach/pdf/tauehuyou_cs-20.pdf"),
        ("農林水産省「2025年農林業センサス（確定値）」第2巻 総括編 表Ⅱ-14",
         "https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&lid=000001487916"),
        ("", None),
        ("注記", None),
        ("・直播栽培面積は都道府県への任意の聞き取り調査で、未回答の県もあり、統計としての精度は確保されていない（農水省注記）。", None),
        ("・2018年以前の値は、上記資料p.3の棒グラフ（ベクター図形）の長さを面積に換算した読取値。R1〜R5の公表値との差は±25ha以内。10ha単位に丸め。", None),
        ("・調査の水稲作付面積は飼料用米・WCS用稲・子実前刈取用稲を含むため、作物統計の水稲（子実用）作付面積とは一致しない。", None),
        ("・令和6年産は乾田・湛水の内訳の確認が一部不十分なため、内訳の和と合計が一致しない（農水省注記）。", None),
        ("・青字＝出典からの転記値、黒字＝数式による計算値。", None),
    ]
    for i, (text, url) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(name=FONT, size=12 if text in ("出典", "注記") else 10, bold=text in ("出典", "注記"))
        if url:
            link = ws.cell(row=i, column=2, value=url)
            link.hyperlink, link.font = url, Font(name=FONT, size=10, color="0563C1", underline="single")
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 70

    out = HERE / "水稲直播栽培_データ.xlsx"
    wb.save(out)
    print("wrote", out)


def main():
    set_font()
    area, region, size = read("chokuha_area.csv"), read("chokuha_region_r6.csv"), read("suito_size_class.csv")
    chart_trend(area, HERE / "chart_chokuha_trend.png")
    chart_index(area, HERE / "chart_ishoku_vs_chokuha.png")
    chart_region(region, HERE / "chart_chokuha_region.png")
    chart_size(size, HERE / "chart_suito_size.png")
    build_excel(area, region, size)


if __name__ == "__main__":
    main()
